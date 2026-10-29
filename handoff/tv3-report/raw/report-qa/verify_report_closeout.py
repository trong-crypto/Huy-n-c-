from pathlib import Path
from hashlib import sha256
from datetime import datetime, timezone, timedelta
from docx import Document
from pypdf import PdfReader
from zipfile import ZipFile
import csv
import json
import re

root = Path(__file__).resolve().parents[3]
report = root/'report/Nhom12_HuyenCo_Report.docx'
qa = Path(__file__).parent
pdf = qa/'Nhom12_HuyenCo_Report.pdf'
digest = lambda p: sha256(p.read_bytes()).hexdigest()
doc = Document(report)
pages = PdfReader(pdf).pages
assert len(pages) == 19 and len(doc.tables) == 11 and len(doc.inline_shapes) == 6
assert [i+1 for i, p in enumerate(pages) if float(p.mediabox.width) > float(p.mediabox.height)] == list(range(15, 20))
assert all(len(p.extract_text().strip()) > 80 for p in pages)
with ZipFile(report) as archive:
    assert archive.testzip() is None
metadata = json.loads((root/'report/authors.json').read_text(encoding='utf-8'))
alltext = '\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for member in metadata['members']:
    assert member['name'] in alltext and member['studentId'] in alltext
    assert member['name'] in doc.core_properties.author
assert metadata['groupNumber'] == 12 and len(metadata['members']) == 3
assert not re.search(r'\b(?:TODO|TBD|PLACEHOLDER)\b|chatgpt-content-reference|::chatgpt', alltext, re.I)
assert not doc.element.xpath('.//w:pBdr') and not doc.styles.element.xpath('.//w:pBdr')
assert all(t.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for t in doc.tables)
assert len(doc.element.xpath('.//wp:docPr')) == 6
assert all(n.get('descr') for n in doc.element.xpath('.//wp:docPr'))
headings = [int(m.group(1)) for p in doc.paragraphs if (m := re.match(r'^(\d+)(?:\s|\.)', p.text)) and p.style.style_id == 'Heading1']
assert headings == list(range(1, 16)), headings
refs = [p.text for p in doc.paragraphs if re.match(r'^\[\d+\] ', p.text)]
assert [int(re.match(r'^\[(\d+)\]', r)[1]) for r in refs] == list(range(1, 14))
body = '\n'.join(p.text for p in doc.paragraphs if p.text not in refs)
cited = set()
for block in re.findall(r'\[([\d,\s\u2013-]+)\]', body):
    for part in block.split(','):
        ns = [int(n) for n in re.findall(r'\d+', part)]
        if len(ns) == 1:
            cited.add(ns[0])
        elif len(ns) == 2:
            cited.update(range(ns[0], ns[1]+1))
assert set(range(1, 14)) <= cited, cited
parse = lambda v: float(v.replace('.', '').replace(',', '.'))
aggregate = list(csv.DictReader((root/'docs/assets/x01/aggregate.csv').open(encoding='utf-8-sig')))
runs = list(csv.DictReader((root/'docs/assets/x01/per-run.csv').open(encoding='utf-8-sig')))
bygroup = {(int(x['clients']), x['mode']): x for x in aggregate}
for row in doc.tables[5].rows[1:]:
    vals = [c.text for c in row.cells]
    src = bygroup[(int(vals[0]), vals[1])]
    assert int(parse(vals[2])) == int(src['cohortRequests'])
    for cell, key in zip(vals[3:], ['medianMs_median', 'p95Ms_median', 'p99Ms_median', 'throughputPerSecond_median']):
        assert abs(parse(cell)-float(src[key])) <= 0.0050001, (cell, key)
for row in doc.tables[6].rows[1:]:
    vals = [c.text for c in row.cells]
    src = bygroup[(int(vals[0]), vals[1])]
    for cell, key in zip(vals[2:], ['p95Ms', 'throughputPerSecond']):
        bounds = cell.split('\u2013')
        assert len(bounds) == 2
        for val, suffix in zip(bounds, ['min', 'max']):
            assert abs(parse(val)-float(src[key+'_'+suffix])) <= 0.0050001
byrun = {(int(x['repeat']), int(x['clients']), x['mode']): x for x in runs}
seen = set()
for row in doc.tables[10].rows[1:]:
    vals = [c.text for c in row.cells]
    key = (int(vals[0]), int(vals[1]), vals[2])
    src = byrun[key]
    seen.add(key)
    assert int(parse(vals[3])) == int(src['cohortRequests'])
    for cell, metric in zip(vals[4:], ['medianMs', 'p95Ms', 'throughputPerSecond']):
        assert abs(parse(cell)-float(src[metric])) <= 0.0050001
assert len(seen) == 18 and sum(int(x['cohortRequests']) for x in runs) == 46027
assert sum(int(x['outOfWindowSends']) for x in runs) == 1
assert '0,022' in alltext and '46.027' in alltext and '59' in alltext
pub = json.loads((root/'docs/assets/x01/publication-manifest.json').read_text())
for name, hash_value in pub['files'].items():
    assert digest(root/'docs/assets/x01'/name) == hash_value
assert digest(root/'docs/53-X01-BENCHMARK-ANALYSIS.md') == pub['reportSha256']
assert all(digest(root/p) == h for p, h in pub['sourcePins'].items())
analysisdir = root/'scripts/generated/x01_analysis_20261029_112956_819906'
inputs = json.loads((analysisdir/'input_checksums.json').read_text())
for group, checkpoint in [('matrix', 'x01_matrix_completed.json'), ('reliability', 'x01_reliability_completed.json')]:
    target = json.loads((root/'scripts/generated'/checkpoint).read_text())
    for relative, hash_value in inputs[group].items():
        assert digest(root/target['evidenceDirectory']/relative) == hash_value, relative
new = qa/'pages-shipping'
old = qa/'pages-reviewed'
changed = [n for n in range(1, 20) if digest(new/f'page-{n}.png') != digest(old/f'page-{n}.png')]
assert changed == [7, 15, 16, 17, 18, 19]
verification = {
    'status': 'PASS', 'checkedAt': datetime.now(timezone(timedelta(hours=7))).isoformat(),
    'pages': 19, 'landscapePages': list(range(15, 20)), 'tables': 11, 'drawingsWithAltText': 6,
    'numberedSections': 15, 'referencesPresentAndCited': 13, 'groupNumber': 12, 'authorsChecked': 3,
    'aggregateRowsMatched': 6, 'perRunRowsMatched': 18, 'rangeRowsMatched': 6,
    'cohortRequests': 46027, 'outOfWindowSends': 1, 'originalPublicationAndToolSourcePinsUnchanged': True,
    'originalMatrixFilesChecked': len(inputs['matrix']), 'originalReliabilityFilesChecked': len(inputs['reliability']),
    'allPagesVisuallyReviewed': True,
    'reviewMethod': 'Pages 1-19 reviewed; final pages 7 and 15-19 re-inspected; other 13 final page PNGs byte-identical to reviewed version',
    'reportSha256': digest(report), 'qaPdfSha256': digest(pdf),
    'pagePngSha256': {f'page-{n}.png': digest(new/f'page-{n}.png') for n in range(1, 20)}
}
indexpath = root/'report/evidence-index.json'
index = json.loads(indexpath.read_text(encoding='utf-8'))
for path, hash_value in index['sourceFiles'].items():
    assert digest(root/path) == hash_value, path
for name, hash_value in index['figures'].items():
    assert digest(root/'report/figures'/name) == hash_value, name
index.update({
    'status': 'COMPLETE_RENDERED_REVIEWED', 'completedAt': verification['checkedAt'],
    'reportSha256': digest(report), 'pageCount': 19, 'authors': metadata,
    'evidencePackage': 'report/evidence/X01-Evidence.zip', 'verification': verification,
    'compilatioRun': False, 'submitted': False
})
(qa/'verification.json').write_text(json.dumps(verification, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
indexpath.write_text(json.dumps(index, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
print(json.dumps({k: v for k, v in verification.items() if k not in ['pagePngSha256', 'reviewMethod']}, ensure_ascii=True))
