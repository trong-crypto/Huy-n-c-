"""Build the Huyen Co technical report from reviewed X01 data using python-docx.

Use the bundled document Python runtime. No database, network, Git or product
source mutation. Group/authors are taken only from report/authors.json.
"""
from __future__ import annotations
from datetime import datetime, timezone, timedelta
import csv
import hashlib
import json
import os
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
VN = timezone(timedelta(hours=7))


def read(path): return json.loads(path.read_text(encoding="utf-8-sig"))
def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def num(value, digits=2): return f"{value:,.{digits}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def load_evidence():
    checkpoint = read(ROOT / "scripts/generated/x01_analysis_completed.json")
    if checkpoint["status"] != "COMPLETE": raise ValueError("X01.6 must be reviewed before writing the paper")
    analysis = read(ROOT / checkpoint["analysisDirectory"] / "analysis.json")
    publication = read(ROOT / "docs/assets/x01/publication-manifest.json")
    if publication["status"] != "REVIEWED": raise ValueError("Figures have not passed visual review")
    for name, sha in publication["files"].items():
        if digest(ROOT / "docs/assets/x01" / name) != sha: raise ValueError("Published X01 artifact changed: " + name)
    if digest(ROOT / "docs/53-X01-BENCHMARK-ANALYSIS.md") != publication["reportSha256"]:
        raise ValueError("Analysis prose changed after review")
    if analysis["totals"]["measurement"] != 46027 or analysis["totals"]["outOfWindowSends"] != 1:
        raise ValueError("This report must be deliberately reviewed for a different matrix")
    return checkpoint, analysis, publication


def architecture(path):
    im = Image.new("RGB", (2200, 1100), "white")
    draw = ImageDraw.Draw(im)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 38)
    small = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 32)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 43)
    def box(bounds, title, body, fill="#F0F5F9"):
        draw.rounded_rectangle(bounds, radius=16, fill=fill, outline="#506577", width=3)
        x0,y0,x1,y1 = bounds
        draw.text((x0+22,y0+20), title, font=bold, fill="#152B3B")
        draw.multiline_text((x0+22,y0+86), body, font=regular, fill="#263746", spacing=12)
    def arrow(start,end,label=None):
        draw.line((start,end), fill="#2F637C", width=6)
        x,y=end; sx,sy=start
        if y==sy:
            direction=1 if x>sx else -1
            draw.polygon([(x,y),(x-20*direction,y-13),(x-20*direction,y+13)],fill="#2F637C")
        else:
            direction=1 if y>sy else -1
            draw.polygon([(x,y),(x-13,y-20*direction),(x+13,y-20*direction)],fill="#2F637C")
        if label:
            pos=((sx+x)//2-85,(sy+y)//2-48)
            bounds=draw.textbbox(pos,label,font=small)
            draw.rectangle((bounds[0]-8,bounds[1]-5,bounds[2]+8,bounds[3]+5),fill="white")
            draw.text(pos,label,font=small,fill="#263746")
    box((40,80,630,330),"Client Next.js","Form và kết quả\nPhiên cookie và lịch sử\nSTOMP SUBSCRIBE rồi SEND")
    box((800,80,1390,330),"Lớp giao tiếp Spring","HTTP REST và WebSocket /ws\nValidator và Origin allowlist\nĐịnh tuyến theo socket")
    box((1570,80,2160,330),"Dispatcher N01","Registry session và requestId\nDirect hoặc worker + queue\nDeadline và terminal replay")
    box((1570,650,2160,940),"Dịch vụ nghiệp vụ","Mai Hoa và Tam Thế\nLịch và tra cứu 64 quẻ\nRegistry nguồn và chính sách")
    box((800,650,1390,940),"PostgreSQL qua JDBC","Tri thức và provenance\nMigration V1 đến V12\nOwner và snapshot lịch sử")
    box((40,650,630,940),"Tiến trình và kết quả","ACCEPTED đến LOOKUP\nCOMPLETED hoặc FAILED\nChỉ gửi về socket sở hữu")
    arrow((630,200),(800,200),"JSON")
    arrow((1390,200),(1570,200))
    arrow((1860,330),(1860,650),"Tính / chờ")
    arrow((1570,790),(1390,790),"JDBC")
    arrow((1390,875),(1570,875))
    draw.line([(1860,940),(1860,982),(330,982),(330,955)],fill="#2F637C",width=6)
    arrow((330,955),(330,940))
    arrow((330,650),(330,330))
    draw.text((55,1000),"REST phục vụ lịch, tra quẻ và lịch sử; STOMP phục vụ tính toán và progress.\nCookie owner dùng cho lịch sử; khóa idempotency dùng session STOMP hiện tại.",font=small,fill="#405261",spacing=8)
    im.save(path)


def content(analysis, authors):
    blocks=[]
    def h(text,level=1): blocks.append({"kind":"heading","text":text,"level":level})
    def p(text): blocks.append({"kind":"paragraph","text":text})
    def table(caption,headers,rows,widths=None): blocks.append({"kind":"table","caption":caption,"headers":headers,"rows":rows,"widths":widths})
    def page(): blocks.append({"kind":"page"})
    def figure(path,caption): blocks.append({"kind":"figure","path":path,"caption":caption})
    blocks.append({"kind":"title","text":"HUYỀN CƠ"})
    blocks.append({"kind":"subtitle","text":"Ứng dụng tra cứu tri thức qua mạng\nvới điều phối hữu hạn và truy vết nguồn"})
    names=authors.get("members",[])
    author_text="\n".join(m["name"]+(" — "+m["studentId"] if m.get("studentId") else "") for m in names) or "Nhóm phát triển Huyền Cơ"
    p((f"Nhóm {authors['groupNumber']} · " if authors.get("groupNumber") else "")+author_text)
    p("Bài tập lớn Lập trình mạng · Năm học 2026–2027 · Ngày 29 tháng 10 năm 2026")
    h("Tóm tắt",2)
    p("Huyền Cơ là ứng dụng Client–Server tra cứu tri thức Mai Hoa Dịch Số và Diễn Cầm Tam Thế. Bài toán kỹ thuật là tiếp nhận nhiều yêu cầu tính toán qua mạng, trả tiến trình đúng phiên, kiểm soát công việc trùng và quá tải, đồng thời bảo toàn căn cứ của dữ liệu số hóa. Hệ thống sử dụng Next.js, Spring Boot, STOMP trên WebSocket và PostgreSQL; Server chịu trách nhiệm tính toán và kiểm chính sách nguồn. Đóng góp triển khai gồm dispatcher với worker/hàng đợi hữu hạn, registry idempotency theo phiên và kết quả có provenance. Thực nghiệm trên cùng máy sử dụng 2/5/20 Client, hai chế độ direct/bounded và ba lần lặp, tạo 18 lượt đo. Cohort theo timestamp có 46.027 yêu cầu thành công, không lỗi ở mức tải đã chọn. Với 20 Client, median độ trễ của bounded là 5,51 ms so với 6,47 ms của direct, nhưng median P99 cao hơn, 36,41 ms so với 22,90 ms. Kiểm độ tin cậy riêng đạt bảy tình huống; stress 2 worker/queue 4 thể hiện từ chối có kiểm soát và phục hồi. Một SEND muộn 0,022 ms được công bố; số đo không được diễn giải thành công suất cực đại hay ưu thế trên mọi mức tải. [5–7]")
    p("Từ khóa: Client–Server; WebSocket; STOMP; hàng đợi hữu hạn; idempotency; truy vết nguồn.")
    h("1 Giới thiệu")
    p("Một ứng dụng tra cứu tri thức có thể trả kết quả đúng ở một phiên nhưng vẫn thất bại khi nhiều Client gửi đồng thời, gửi lặp hoặc rời kết nối. Vì vậy mục tiêu của Huyền Cơ gồm cả nghiệp vụ lẫn thiết kế giao tiếp: request phải được kiểm tra, tiến trình được tương quan bằng requestId, kết quả trả đúng socket và giới hạn tài nguyên có hành vi rõ ràng. Nguồn sách còn có chỗ mờ hoặc khác cách áp dụng, nên hệ thống phải phân biệt nội dung đã đối chiếu với quy ước nhóm. [8–10]")
    p("Báo cáo trình bày kiến trúc, protocol, thuật toán ở phạm vi đã triển khai và bằng chứng thực nghiệm. Trọng tâm đánh giá là lát cắt LINES trên dữ liệu TEST cố định; các chức năng TIME, OBJECT_COUNT, Tam Thế, lịch và lịch sử được mô tả theo module và biên bản nghiệm thu tương ứng. Nội dung văn hóa truyền thống được trình bày có nguồn và giới hạn, không coi là dự đoán khoa học.")
    page()
    h("2 Cơ sở kỹ thuật và cách tiếp cận liên quan")
    p("WebSocket cung cấp kênh giao tiếp hai chiều sau opening handshake trên TCP. STOMP bổ sung frame với lệnh, header và body; Client sử dụng SEND và SUBSCRIBE để gửi yêu cầu, nhận thông điệp theo destination. Trong Spring, handler ứng dụng và broker nội bộ đảm nhiệm các phần định tuyến khác nhau. Các chuẩn này cung cấp cơ chế vận chuyển; hợp đồng nghiệp vụ và nghĩa của các queue vẫn do ứng dụng định nghĩa. [1–3]")
    p("Ở lớp thực thi, ThreadPoolExecutor cho phép đặt số luồng, hàng đợi và chính sách khi không nhận thêm công việc. Nhóm dùng fixed-size pool, ArrayBlockingQueue và AbortPolicy. Đây là áp dụng cơ chế thư viện để giới hạn việc tính toán, không phải thuật toán thread pool mới. Baseline direct thực hiện phép tính trên luồng gọi dispatcher của Spring; direct vẫn có thể phục vụ nhiều Client đồng thời. [4,8]")
    p("Hai mode chia sẻ thuật toán, cơ chế xác thực input, registry idempotency và dữ liệu nguồn. So sánh này chỉ thay điều phối; không so một baseline trả lỗi hoặc tính khác nghiệp vụ với một bản trả kết quả đầy đủ. Báo cáo không có thí nghiệm đối đầu với sản phẩm thương mại hoặc khảo sát đầy đủ công trình học thuật tương tự, nên không khẳng định tính mới ở cấp độ nghiên cứu thế giới.")
    h("3 Định nghĩa bài toán")
    table("Bảng 1. Đầu vào và đầu ra theo nhóm chức năng",["Nhóm","Đầu vào","Đầu ra"],[
        ["Mai Hoa","LINES hoặc số vật/giờ Chi hoặc thời điểm","Quẻ chủ/hỗ/biến, hào động, bước tính; lời luận khi chọn reading"],
        ["Tam Thế","Ngày/giờ sinh, loại lịch, tháng nhuận, category","Các mục luận có nguồn; trạng thái không khớp nguồn được giữ rõ"],
        ["Tra cứu và lịch","Tên/mã quẻ; ngày và loại lịch","Thông tin quẻ hoặc lịch chuyển đổi/Can Chi"],
        ["Lịch sử","Phiên owner; requestId kết quả đã tính","Snapshot kết quả, tìm/xem/xóa đúng owner"]],[2.6,5.6,8.4])
    p("Đầu vào được kiểm trên Server; Client không được cấp quyền duyệt nguồn hoặc tự gửi snapshot để lưu. Miền lịch dự án là 1900–2099 với UTC+7 cố định. Các yêu cầu tính toán mang UUID; kết nối STOMP được giữ riêng với cookie owner dùng cho lịch sử. Khi không đủ dữ liệu hoặc điều kiện chính sách, Server trả lỗi có mã thay vì tạo nội dung bù. [9–12]")
    page()
    h("4 Giải pháp đề xuất")
    p("Yêu cầu qua STOMP được parse JSON/UTF-8 nghiêm ngặt, sau đó dispatcher kiểm registry theo cặp sessionId/requestId. ID mới được nhận vào vòng đời Run với deadline; ID trùng có cùng fingerprint chia sẻ công việc đang chạy hoặc replay terminal còn hiệu lực. Nội dung khác cùng ID bị từ chối. Trong bounded, chỉ worker thực thi nghiệp vụ; queue đầy trả SERVER_BUSY, không chuyển việc sang luồng nhận STOMP để né giới hạn. [8]")
    p("Service tra registry nguồn trước khi tính hoặc lấy lời luận. Dữ liệu kết quả giữ ruleVersion, ruleBasis, sourceStatus, sourceReferences và decisionReferences. Chính sách SOURCE_VERIFIED và TEAM_DECISION được phân biệt với trạng thái nguồn. TIME dùng quy ước chia 6 đã chốt cho dự án; OBJECT_COUNT giữ TEAM_DECISION/SOURCE_AMBIGUOUS theo hồ sơ. Một quyết định nhóm không biến bản in chưa rõ thành chứng cứ ấn bản đã được xác minh. [9–10]")
    h("5 Kiến trúc hệ thống")
    figure("figures/architecture.png","Hình 1. Các thành phần giao tiếp và xử lý của Huyền Cơ. Mũi tên mô tả luồng logic; truy cập dữ liệu và trả thông điệp do các module tương ứng thực hiện.")
    p("Client quản lý nhập liệu và hiển thị; Server giữ thuật toán và quyền truy cập. REST phục vụ các thao tác tra cứu/lịch/lịch sử; WebSocket duy trì kênh progress và kết quả tính toán. PostgreSQL chứa tri thức có nguồn, metadata review và snapshot lịch sử. H01 lưu định danh owner bằng hash token, tách khỏi session STOMP dùng cho định tuyến và N01. [11]")
    page()
    h("6 Thiết kế giao tiếp mạng")
    h("Thiết lập kết nối và trình tự thông điệp",2)
    p("Client mở WebSocket /ws, hoàn tất STOMP CONNECT/CONNECTED rồi subscribe /user/queue/progress, /user/queue/results và /user/queue/errors trước SEND. Tuyến tính toán là /app/mai-hoa/calculate hoặc /app/tam-the/calculate. Với lịch sử, Client bootstrap cookie owner qua POST /api/v1/session trước handshake; cookie hợp lệ được gắn owner phía Server. Client công khai không cookie vẫn tính được nhưng không có snapshot owned để lưu. [9,11]")
    blocks.append({"kind":"code","text":'{\n  "requestId": "550e8400-e29b-41d4-a716-446655440002",\n  "type": "MAI_HOA",\n  "payload": {\n    "method": "LINES",\n    "linesBottomUp": [true, false, true, false, true, false],\n    "movingLine": 2\n  }\n}'})
    p("Progress mang requestId, type, sequence, step và message; chuỗi thành công gồm ACCEPTED, VALIDATING, CALCULATING, LOOKUP. Terminal mang requestId/type/status/data/error; thành công COMPLETED có error=null, thất bại FAILED có mã lỗi. Các frame được trả về socket tương ứng; Client không được SEND trực tiếp vào broker queue hoặc SUBSCRIBE queue của phiên khác. Hợp đồng JSON v1.2 khác với phiên bản STOMP 1.2 dù cùng con số. [2,9]")
    h("Lỗi và đóng kết nối",2)
    p("JSON hỏng, field lạ và input sai bị từ chối; xung đột requestId trả INVALID_INPUT; từ chối admission trả SERVER_BUSY. Deadline của Run gồm thời gian chờ queue; vượt hạn trả PROCESSING_TIMEOUT và chặn terminal/progress muộn. Disconnect hủy Run/timer, yêu cầu interruption và xóa registry của phiên; reconnect tạo session mới. Interruption là best effort, không bảo đảm JDBC đang kẹt dừng ngay. [8]")
    p("Ghi/xóa lịch sử yêu cầu cookie owner, custom header và CSRF hợp lệ; truy cập history của owner khác trả 404. Origin allowlist và bảo vệ cookie/CSRF là lớp riêng với session routing. Các lượt benchmark dùng loopback HTTP/WS; chưa phải nghiệm thu TLS hoặc triển khai Internet. [11]")
    page()
    h("7 Triển khai")
    table("Bảng 2. Thành phần chính và trách nhiệm",["Module","Vai trò"],[
        ["CalculationController","Parse nghiêm ngặt, kiểm hợp đồng, gọi service và ánh xạ lỗi"],
        ["CalculationDispatcher","Admission nguyên tử, worker/queue, fingerprint và replay terminal"],
        ["CalculationSessions","Run, progress/terminal theo socket, deadline và cleanup"],
        ["MaiHoaService / TamTheService","Thuật toán và bảng tra có điều kiện nguồn"],
        ["RuleRegistry / JDBC repositories","Cổng chính sách, snapshot dữ liệu và truy vết nguồn"],
        ["H01 và Client","Owner/history/CSRF; nhập liệu, nhận STOMP và hiển thị"]],[5.4,11.2])
    h("Thuật toán và dữ liệu",2)
    p("Với LINES, sáu hào được đánh số từ dưới lên. Hỗ hạ lấy hào 2–4, hỗ thượng lấy hào 3–5; quẻ biến đảo duy nhất hào động. OBJECT_COUNT dùng số vật lấy thượng quái, chỉ số giờ Chi lấy hạ quái; tổng số gốc chia 6 lấy hào, dư 0 quy về 6. TIME chuẩn hóa UTC+7 và lấy năm/tháng/ngày âm cùng giờ Chi theo quy ước đã công bố. Luận chi tiết là chức năng opt-in riêng với mẫu đã review, không phải output mặc định của workload benchmark. [9–10]")
    p("Tam Thế chuẩn hóa lịch, lấy Can Chi năm âm và hành mệnh rồi tra bảng theo loại nội dung. EXTENDED giữ Lộc, Sang hèn và Điền sản thành các mục riêng; không có giao điểm nguồn của Lộc được ghi NO_SOURCE_MATCH, không tự sinh lời thay thế. Lịch C01 dùng dataset cố định tạo và cross-check với Time4J 5.9.4; kiểm cùng thư viện không được gọi là nguồn thiên văn độc lập. [12]")
    p("Schema V1–V12 phân biệt sách/mục nguồn, bát quái/quẻ/hào, quy tắc, nội dung tra và metadata review; V11 bổ sung owner/history, V12 dữ liệu mở rộng. H01 lưu snapshot Result DTO đã tính trên Server theo owner; không nhận payload/snapshot hoặc trạng thái duyệt do Client tự cung cấp. Phạm vi VERIFIED được xác định theo từng hồ sơ; tên quẻ đã duyệt không chứng minh mọi thuật toán hay mọi lời luận đã được duyệt. [10–12]")
    p("Các biên bản Client 8E/8F ghi nhận nghiệm thu UI/HTTP/STOMP ở những ca đại diện và hai owner. Chúng được giữ riêng với kết quả Backend và X01; không suy từ vài ca UI thành kiểm hết miền input hoặc mọi tình huống lỗi mạng. [13]")
    page()
    h("8 Thiết lập thực nghiệm")
    h("Phần cứng phần mềm và mạng",2)
    table("Bảng 3. Môi trường đã ghi nhận của ma trận",["Thành phần","Giá trị"],[
        ["CPU","Intel Core i7-13620H; 10 core, 16 logical processor"],
        ["RAM hệ thống","25.458.917.376 byte theo snapshot môi trường"],
        ["Hệ điều hành","Windows 11 Home Single Language 64-bit, 10.0.26300"],
        ["Backend","JDK 17.0.12; Spring Boot 3.5.16; heap 256–512 MiB"],
        ["Runner","Node.js 24.21.0; @stomp/stompjs 7.5.0; ws 8.22.0"],
        ["Database","PostgreSQL 18.6, database benchmark TEST riêng, schema V12"],
        ["Mạng","Cùng máy, loopback 127.0.0.1; HTTP 18091, WS /ws"],
        ["Origin / kết nối","http://127.0.0.1:3173; 20/5/2 socket thật tùy lượt"]],[3.3,13.3])
    table("Bảng 4. Các tham số giữ cố định ngoài mode",["Tham số","Cấu hình"],[
        ["Ma trận","2/5/20 Client × direct/bounded × 3 lần = 18 lượt"],
        ["Thời gian","Warm-up 10 s; đo 60 s; drain tối đa 15 s; nghỉ giữa lượt 5 s"],
        ["Tải Client","Một request chưa terminal; nghỉ ít nhất 200 ms sau terminal"],
        ["Worker / queue","Bounded 4/100; direct không qua queue riêng này"],
        ["Deadline / TTL","10.000 ms gồm queue; idempotency 60 s, tối đa 10.000 ID"],
        ["Giới hạn Run","1.024 phiên; 1.000 toàn Server; 16/phiên"],
        ["JDBC / Flyway","Pool 8, read-only; Flyway tắt trong mọi lượt đo"],
        ["Retry / cookie","Không auto retry/reconnect, không cookie hoặc ghi history"]],[3.3,13.3])
    p("Ma trận chạy ngày 29/10/2026 từ 02:20:47 đến 02:44:49 UTC+7. Mỗi lượt có JVM mới, cùng JAR đã pin; không xóa cache PostgreSQL/OS giữa lượt. Thứ tự cân bằng được ghi tại Phụ lục A; không tự thay thế hoặc chạy lại một lượt để chọn kết quả. [5]")
    page()
    h("Dữ liệu workload và mốc đo",2)
    p("Database huyenco_benchmark_test_20261029_010531 có snapshot tri thức/schema khớp baseline trước và sau mỗi lượt. JAR SHA-256 là 71a9d31d11adda8f50330398e87ccad6029fe7782b0a5b6239eaabf1bad0ab39. Các snapshot runtime có anonymous_sessions=0 và interpretation_history=0. Thuật toán và dữ liệu không đổi giữa hai mode. [5,7]")
    p("Workload là Mai Hoa LINES không reading. Chỉ số Client i tính từ 0, sequence s bắt đầu từ 0 ở mỗi phase; pattern=(17i+s) mod 64; movingLine=1+((i+s) mod 6). Seed cố định, UUID mới mỗi SEND; chu kỳ đo phủ 192 cặp inputCase mỗi lượt, không phải toàn bộ 384 tổ hợp chức năng. Mỗi kết quả được oracle kiểm ba quẻ, hào động, rule/source, requestId, session và chuỗi progress 1–4. [5,7]")
    table("Bảng 5. Định nghĩa chỉ số",["Chỉ số","Cách tính"],[
        ["Cohort","SEND trong [measurement.start, measurement.end)"],
        ["Latency","Terminal nhận tại runner trừ SEND, cùng đồng hồ đơn điệu"],
        ["Median / P95 / P99","COMPLETED hợp lệ; percentile nearest-rank theo từng lượt"],
        ["Throughput","Số hoàn tất thuộc cohort nhận trước end / 60 giây"],
        ["Tổng hợp ba lượt","Median và min–max của ba chỉ số; không phải CI"],
        ["CPU / bộ nhớ","Mẫu khoảng 1 s; CPU % một core; peak lấy mẫu và phạm vi riêng"]],[3.3,13.3])
    h("Sai lệch ranh thời gian và xử lý",2)
    p("Raw có 46.028 dòng gắn nhãn measurement, nhưng SEND tại 20 Client/bounded/lần 1 muộn 0,022 ms sau ranh 60 s. Dòng này vẫn được giữ nguyên. Quy tắc cohort [start,end) áp dụng đồng nhất theo SEND timestamp tạo 46.027 mẫu hợp lệ; median của lượt bị ảnh hưởng sửa 5,5139 thành 5,5143 ms, P95 và throughput không đổi. Đây là sai lệch kiểm soát lịch gửi, không được gọi là strict-window PASS tuyệt đối. [5]")
    p("Sau ma trận đã thêm chốt deadline trước publish và kiểm hồi quy riêng. Bản sửa không được đưa ngược vào 18 lượt cũ; raw/summary gốc và snapshot mã đã đo được giữ. Nếu tiêu chí chấm yêu cầu zero SEND ngoài ranh ở bản chạy gốc, cần một ma trận mới. [5,7]")
    page()
    h("9 Kết quả")
    h("Hiệu năng theo cấu hình",2)
    perf=[]
    for g in analysis["groups"]:
        m=g["metrics"]
        perf.append([str(g["clients"]),g["mode"],num(g["cohortRequests"],0),num(m["medianMs"]["median"]),num(m["p95Ms"]["median"]),num(m["p99Ms"]["median"]),num(m["throughputPerSecond"]["median"])])
    table("Bảng 6. Median của ba chỉ số theo lượt cho mỗi ô",["Client","Mode","Cohort","Med ms","P95 ms","P99 ms","Req/s"],perf,[1.3,2.2,2.3,2.5,2.5,2.5,2.8])
    p("Cả 46.027 yêu cầu cohort hoàn tất hợp lệ; không có lỗi hoặc terminal cohort về sau cửa sổ. Tổng raw gồm 162 readiness, 7.640 warm-up, 46.027 cohort và một SEND muộn, thành 53.830 COMPLETED/215.320 progress. Không tính replay hoặc lỗi từ chối nhanh của bài độ tin cậy vào số đo này. Các hình 2–6 trình bày toàn bộ biến thiên theo lượt, ECDF, tài nguyên và diễn biến theo giây. [5–7]")
    table("Bảng 7. Khoảng min–max của P95 và throughput qua ba lượt",["Client","Mode","P95 min–max ms","Throughput min–max req/s"],[
        [str(g["clients"]),g["mode"],num(g["metrics"]["p95Ms"]["min"])+"–"+num(g["metrics"]["p95Ms"]["max"]),num(g["metrics"]["throughputPerSecond"]["min"])+"–"+num(g["metrics"]["throughputPerSecond"]["max"])] for g in analysis["groups"]],[1.5,2.5,5.5,7.1])
    p("Ở 20 Client, median latency bounded giảm khoảng 14,71%, P95 giảm 1,96% và throughput tăng 0,39% so với direct. Tuy nhiên median P99 tăng khoảng 59%, từ 22,90 lên 36,41 ms; max quan sát 60,75 ms so với 52,75 ms. Các khoảng đo vẫn chồng lấn; kết quả không cho phép kết luận bounded tốt hơn toàn phân bố. [7]")
    p("Hình 3 so throughput với trần nhịp gửi 5 request/s/Client từ khoảng nghỉ 200 ms; Hình 6 giữ cả ba chuỗi theo bin một giây. Throughput đo được khoảng 94–95 request/s ở 20 Client không phải công suất cực đại Server. [7]")
    page()
    h("Độ tin cậy và tài nguyên",2)
    table("Bảng 8. Bảy tình huống qua STOMP và PostgreSQL thật",["Tình huống","Kết quả chính"],[
        ["In-flight trùng / xung đột","Một progress chain/kết quả gốc; đảo key và UUID hoa được dedupe; khác payload INVALID_INPUT"],
        ["Terminal replay / xung đột","Replay đúng terminal, không progress; conflict không thay kết quả gốc"],
        ["20 phiên cùng ID","20 input khác nhận đúng kết quả/progress của phiên tương ứng"],
        ["Reconnect","Socket mới gửi cùng ID/input tạo chuỗi tính mới"],
        ["TTL mặc định 60 s","Gửi lại sau 61,506 s tạo 4 progress và kết quả mới"],
        ["Burst tự nhiên 20 × 8","10 COMPLETED, 150 SERVER_BUSY; ID cũ bị từ chối và ID mới phục hồi"],
        ["Queue 2 worker / 4 việc","Có giữ khóa TEST: 6 nhận, 154 SERVER_BUSY; sau mở khóa cả 6 hoàn thành"]],[4.8,11.8])
    p("Suite stress giữ các giới hạn còn lại lớn hơn burst: tám việc/phiên dưới 16, tổng 160 dưới 1.000, registry dưới 10.000. Hai JDBC waiter được quan sát khi giữ khóa public.hexagrams; 154 từ chối về trước mở khóa, sau đó sáu việc hoàn tất. Điều kiện này cùng mã dispatcher/cấu hình loại trừ các giới hạn admission khác trong ca đã kiểm. Khóa thử chỉ trên database benchmark, giao dịch rollback và watchdog 4 s; không sửa dữ liệu hoặc thêm sleep vào nghiệp vụ. [6]")
    p("CPU Node trung bình mẫu trong cửa sổ khoảng 1,28–7,16% một logical core; khoảng cao nhất 21,81%. Peak RSS Node khoảng 170,65 MiB; peak working set Java khoảng 402,27 MiB. Java CPU/RAM/OS thread lấy mẫu toàn vòng đời gồm startup/JIT/warm-up, không phải steady state 60 s. Node dùng đoạn sample thực 58,23–58,37 s trong cửa sổ. Không cộng hai CPU khác phạm vi; PostgreSQL CPU không đo riêng. [7]")
    p("Cả 18 cặp snapshot trước/sau và kiểm chỉ đọc X01.6 lúc 11:21:28 UTC+7 ngày 29/10 đều khớp 16 bảng tri thức/migration, schema và runtime trống. Trong X01.5, hai Java/Node thử đã dừng và các khóa được giải phóng; không ghi vào huyenco hoặc TEST nguồn để chạy thí nghiệm. [5–7]")
    page()
    h("10 Thảo luận")
    h("Ý nghĩa của độ trễ và nhịp gửi",2)
    p("Ở 2 và 5 Client, thay đổi median/P95 nhỏ và không đồng nhất. Khi tăng lên 20 Client, median khoảng 5–6 ms và P95 khoảng 12 ms phản ánh nhiều yêu cầu cùng hoạt động hơn, nhưng thí nghiệm chưa tách thời gian queue, JDBC, GC hay scheduling. Không thể quy toàn bộ chênh lệch cho worker pool. Ba lượt/ô và các khoảng chồng lấn chưa hỗ trợ khẳng định ý nghĩa thống kê; min–max không phải confidence interval. [7]")
    p("P99 của bounded ở 20 Client cao hơn là một bất lợi cần giữ trong kết luận. Median thấp hơn có thể đi cùng vài request ở phần đuôi chậm hơn; ECDF ở Hình 4 cho thấy vì sao không thể thay mô tả phân bố bằng một con số trung tâm. Báo cáo không loại outlier để cải thiện kết quả. Không có telemetry đủ để kết luận nguyên nhân cụ thể của đuôi dài. [7]")
    p("Vòng kín với nghỉ ít nhất 200 ms giới hạn nhịp gửi gần 10/25/100 request/s trước khi tính thời gian xử lý. Do đó throughput gần nhau là điều có thể xảy ra ngay cả khi cơ chế điều phối khác; đây là cách đọc phù hợp thiết kế workload, không phải chứng minh hai hệ thống có cùng năng lực cực đại. Benchmark dùng kết nối đã thiết lập; thời gian HTTP upgrade/STOMP CONNECTED được ghi riêng. [5,7]")
    h("Kiểm soát quá tải và cách ly phiên",2)
    p("Giá trị kỹ thuật rõ hơn của bounded trong bộ thử là hành vi có giới hạn: từ chối khi pool/queue đủ chỗ đã bị sử dụng, không crash, không chuyển trả sai phiên và tiếp nhận việc mới sau burst. Bằng chứng đó đến từ suite stress riêng, không từ bảng latency thường. Idempotency cũng được kiểm riêng vì cả direct và bounded đã cùng bật registry; không thể lấy chênh lệch hai mode làm số đo lợi ích riêng của dedupe. [6,8]")
    h("Nguồn dữ liệu và khả năng tái lập",2)
    p("Hash JAR, nguồn runner, schema và dữ liệu giữ được quan hệ giữa cấu hình và kết quả. Dù vậy cùng máy còn chịu cache PostgreSQL/OS và tác vụ nền; JVM mới không làm các lượt hoàn toàn độc lập về trạng thái hệ thống. Tài nguyên một giây/mẫu có thể bỏ lỡ đỉnh ngắn, Java gồm startup và Node event-loop histogram có resolution 20 ms/phạm vi toàn lượt. Sai lệch SEND muộn là giới hạn thật của kiểm soát thời gian, cần công bố ngay cả khi ảnh hưởng số liệu tổng hợp nhỏ. [5,7]")
    page()
    h("11 Đóng góp kỹ thuật")
    table("Bảng 9. Đóng góp ở phạm vi project và bằng chứng",["Đóng góp","Triển khai","Bằng chứng"],[
        ["Điều phối hữu hạn","Pool cố định, queue hữu hạn, AbortPolicy, deadline và cleanup","Mã dispatcher; stress có hai waiter, 6 nhận/154 từ chối và phục hồi [6,8]"],
        ["Idempotency theo socket","Fingerprint JSON; trùng in-flight; replay TTL; conflict; phiên độc lập","Bảy ca live gồm TTL60s/reconnect/20phiên; không suy từ latency benchmark [6]"],
        ["Kết quả truy vết nguồn","Registry tách policy/trạng thái; phiên bản/hash/trang và quyết định được giữ","Hợp đồng1.2; schema và snapshot kiểm toán; SOURCE_VERIFIED/TEAM_DECISION [9–12]"],
        ["Quy trình thực nghiệm tái tính","Pin JAR/dataset, raw theo timestamp, bảng/ảnh và công bố sai lệch","18 lượt, auditor độc lập, CSV và hash; một dòng ngoài ranh được giữ [5–7]"]],[3.4,6.1,7.1])
    p("Các đóng góp là kết quả tích hợp và kiểm chứng trong project. Thread pool, WebSocket, STOMP và cache không được tuyên bố là phát minh mới; hình thức UI không được dùng thay bằng chứng cải tiến kỹ thuật. Chưa có ablation bật/tắt idempotency hoặc so các thuật toán queue khác. [1–4,6–8]")
    h("12 Hạn chế")
    p("Bộ đo chỉ dùng LINES, 192 cặp seed, 2/5/20 Client logic cùng máy và workload có khoảng nghỉ. Chưa đánh giá nhiều máy vật lý, jitter/loss/băng thông mạng, tải mở, workload lời luận chi tiết/TIME/Tam Thế/REST/history hoặc mức tải cực đại. Chưa đo tách queue/JDBC và CPU PostgreSQL; chỉ ba lần/ô, không có p-value hoặc CI.")
    p("Registry idempotency nằm trong bộ nhớ một tiến trình; restart mất cache, reconnect là session mới. Interruption không bảo đảm thư viện đang chạy dừng ngay. X01.5 không bao phủ crash giữa phép tính, registry đầy hoặc mọi tình huống timeout/cancel qua mạng; test N01 trước đó có phạm vi riêng. [6,8]")
    p("Tri thức sách có giới hạn ấn bản và một số quy ước biên tập/nhóm. VERIFIED có phạm vi theo bản ghi và hồ sơ review, không đồng nghĩa kiểm chứng khoa học; các lượt review cùng công cụ không được mô tả là hai chuyên gia độc lập. Dữ liệu không khớp nguồn không được lấp bằng lời luận suy diễn. [10,12]")
    page()
    h("13 Hướng phát triển")
    p("Ưu tiên một ma trận mới với final publish gate đã sửa khi cần đáp ứng strict-window tuyệt đối, giữ cả bộ cũ để đối chiếu. Sau đó mở rộng thực nghiệm sang hai máy thật và workload mix có lời luận chi tiết, tách tải mở/vòng kín, đo queue wait và JDBC time bằng telemetry có chi phí được kiểm soát. Tăng số lần lặp và thiết kế phân tích thống kê trước khi thu dữ liệu mới.")
    p("Về hệ thống, có thể khảo sát hạn tải theo owner/kết nối, metrics phân biệt nguyên nhân SERVER_BUSY và idempotency bền vững nếu cần nhiều instance. Triển khai Internet còn cần nghiệm thu TLS/reverse proxy, cookie, chính sách log và vận hành; không suy từ loopback hiện tại thành sẵn sàng vận hành công cộng. Mở rộng nội dung phải bắt đầu từ khảo sát/review nguồn thay vì sinh nội dung thay thế.")
    h("14 Kết luận")
    p("Huyền Cơ kết hợp tra cứu tri thức có nguồn với thiết kế Client–Server qua REST và STOMP/WebSocket. Server chịu trách nhiệm thuật toán, cổng nguồn, vòng đời request và dữ liệu sở hữu. Phần N01 bổ sung worker/queue hữu hạn cùng idempotency theo session, giúp hành vi khi nhiều Client hoặc quá tải có điều kiện và kết quả kiểm tra rõ ràng.")
    p("Ma trận 18 lượt có 46.027 yêu cầu cohort hoàn tất hợp lệ ở mức tải đã chọn; suite độ tin cậy đạt bảy tình huống. Bounded tại 20 Client có median thấp hơn nhưng P99 cao hơn; throughput hai mode gần nhau và bị ràng buộc bởi khoảng nghỉ. Kết quả vì vậy ủng hộ đánh giá dựa trên nhiều chỉ số và bằng chứng quá tải riêng, thay vì khẳng định ưu thế tốc độ toàn diện. Một SEND muộn 0,022 ms và các giới hạn tài nguyên/mạng được giữ trong báo cáo. [5–7]")
    h("Phân công và trách nhiệm",2)
    table("Bảng 10. Vai trò kỹ thuật của nhóm",["Vai trò","Phần việc"],[
        ["TV1","Client Next.js, form, kết quả/nguồn, REST/STOMP và nghiệm thu UI đại diện"],
        ["TV2","Backend, protocol, vòng đời/concurrency/idempotency, benchmark"],
        ["TV3","Nguồn sách, schema/review, lịch/Tam Thế và kiểm toán dữ liệu"]],[2.3,14.3])
    p("Theo tổ chức hiện tại, chủ dự án đảm nhiệm cả TV2 và TV3. Công cụ AI được dùng hỗ trợ thiết kế, mã nguồn, kiểm thử, phân tích và soạn tài liệu; các thành viên chịu trách nhiệm kiểm tra, hiểu và giải thích phần nộp. Không dùng sự tồn tại của công cụ như chứng cứ xác minh nguồn hoặc thay kết quả thực nghiệm.")
    page()
    h("15 Tài liệu tham khảo")
    references=[
        (1,"I. Fette và A. Melnikov. The WebSocket Protocol. RFC 6455, 2011. https://www.rfc-editor.org/rfc/rfc6455"),
        (2,"STOMP. STOMP Protocol Specification Version 1.2. https://stomp.github.io/stomp-specification-1.2.html"),
        (3,"Spring. Spring Framework 6.2 Reference, WebSocket STOMP Overview. https://docs.spring.io/spring-framework/reference/6.2/web/websocket/stomp/overview.html"),
        (4,"Oracle. Java SE 17 API, ThreadPoolExecutor. https://docs.oracle.com/en/java/javase/17/docs/api/java.base/java/util/concurrent/ThreadPoolExecutor.html"),
        (5,"Nhóm Huyền Cơ. X01.4, ma trận 18 lượt và audit cohort. 29/10/2026. docs/51-X01-BENCHMARK-RESULTS.md; artifact matrix trong Phụ lục A."),
        (6,"Nhóm Huyền Cơ. X01.5, nghiệm thu quá tải/idempotency. 29/10/2026. docs/52-X01-OVERLOAD-IDEMPOTENCY-ACCEPTANCE.md; raw/control transcript và independent audit."),
        (7,"Nhóm Huyền Cơ. X01.6, tổng hợp/phân tích và năm biểu đồ. 29/10/2026. docs/53-X01-BENCHMARK-ANALYSIS.md; docs/assets/x01/per-run.csv và aggregate.csv."),
        (8,"Nhóm Huyền Cơ. Thiết kế N01 và mã CalculationDispatcher/CalculationSessions. docs/36-N01-CONCURRENCY-IDEMPOTENCY-HANDOFF.md; server/src/main/java/com/huyenco/websocket/."),
        (9,"Nhóm Huyền Cơ. Hợp đồng JSON/STOMP v1.2. docs/02-API-CONTRACT.md; client/src/types/contract.ts; DTO Java và docs/examples/."),
        (10,"Mai Hoa Dịch Số, bản scan người dùng cung cấp, 672 trang PDF; chưa xác minh đầy đủ ấn bản in. Hash MH 225e9506f4f7481e56618e3f46738467a64cf6231148c959664a23469faf22de. Hồ sơ thực thi: docs/04-MAIHOA-RULES.md, 28-TIME-OWNER-APPROVAL.md và 41-M03-MAI-HOA-INTERPRETATION-HANDOFF.md."),
        (11,"Nhóm Huyền Cơ. H01, phiên/owner/lịch sử và schema. docs/34-H01-SESSION-HISTORY-HANDOFF.md; migrations V11/V12 và JDBC repositories."),
        (12,"Diễn Cầm Tam Thế Diễn Nghĩa, Dương Công Hầu, hiệu Khương Đức, bản scan 360 trang; thông tin ấn bản chưa xác minh đầy đủ. Hồ sơ dự án: docs/05-TAMTHE-RULES.md; 26-C01-CALENDAR-HANDOFF.md; 42-D04-TAM-THE-EXTENDED-WEALTH-HANDOFF.md."),
        (13,"Nhóm Huyền Cơ. TV1 Phase 8, bằng chứng UI đại diện 8E/8F. docs/44-TV1-PHASE8A-ACCEPTANCE.md. Đây là hồ sơ nghiệm thu có phạm vi, không thay toàn bộ kiểm sản phẩm.")]
    for index,text in references:p(f"[{index}] {text}")
    p("Các tài liệu Web [1–4] được truy cập ngày 29/10/2026. Dẫn trang sách trong kết quả dùng metadata theo từng mục nguồn; không suy một độ lệch trang chung cho cả cuốn.")
    page()
    h("Phụ lục A Số liệu từng lượt và truy vết")
    rows=[]
    for run in analysis["runs"]:
        rows.append([str(run["repeat"]),str(run["clients"]),run["mode"],num(run["cohortRequests"],0),num(run["medianMs"]),num(run["p95Ms"]),num(run["throughputPerSecond"])])
    table("Bảng A1. Giữ nguyên thứ tự chạy, không chọn lại lượt",["Lần","Client","Mode","Cohort","Med ms","P95 ms","Req/s"],rows,[1.2,1.4,2.2,2.6,2.7,2.7,3.3])
    p("Ô 20/bounded/lần 1 có 5.683 request cohort; một dòng measurement muộn 0,022 ms được giữ tại boundary-exclusions.csv. P95 nearest-rank theo lượt; các bảng tổng hợp dùng median của ba chỉ số. Toàn bộ raw/snapshot/log là artifact local đã pin, không tự xuất hiện chỉ vì push source code lên GitHub.")
    p("Gốc raw X01.4: scripts/generated/x01_matrix_20261029_022047_427181/. Raw X01.5: scripts/generated/x01_reliability_20261029_103446_353616/. Bản phân tích dùng cho báo cáo: "+analysis["analysisDirectory"]+"/. Các manifest SHA-256 và report/evidence-index.json liên kết các tệp cần bàn giao.")
    p("Tái tính offline khi có đầy đủ artifact: python scripts/audit_x01_matrix.py; python scripts/audit_x01_reliability.py --evidence scripts/generated/x01_reliability_20261029_103446_353616; python scripts/summarize_x01_benchmark.py. Không đưa mật khẩu, cookie hoặc CSRF vào gói bằng chứng.")
    figures=[
        ("01-latency","Hình 2. Median, P95 và P99 theo từng lượt ở 2/5/20 Client. Dấu lớn là median ba chỉ số; chấm nhỏ là từng lượt; thanh min–max quan sát không phải CI. Direct màu xanh, bounded màu nâu; một SEND ngoài cửa sổ không thuộc cohort."),
        ("02-throughput","Hình 3. Throughput thành công và tỷ lệ so với trần nhịp gửi từ nghỉ 200 ms. Trần 5 × số Client request/s chưa kể xử lý/lập lịch. Dữ liệu không phải công suất cực đại; thanh min–max qua ba lượt."),
        ("03-latency-ecdf","Hình 4. ECDF gộp request của ba lượt theo từng ô. Trục X log, toàn dải độ trễ được giữ; đường ngang ở 95%. ECDF gộp không thay median các P95 theo lượt và không coi từng request là một lượt độc lập."),
        ("04-resources","Hình 5. CPU và bộ nhớ được lấy mẫu của Java/Node. Java gồm startup/JIT/warm-up; Node chỉ sample trong cửa sổ. CPU là % một logical core, working set/RSS là MiB và peak được lấy mẫu; không phải heap used, không đo CPU PostgreSQL riêng."),
        ("05-throughput-timeline","Hình 6. Số kết quả thuộc cohort hoàn tất trong từng bin một giây ở 20 Client. Giữ cả ba lần, không làm mượt; không gồm readiness/warm-up/SEND muộn. Số đếm theo giây không chứng minh queue đầy hoặc giải thích nguyên nhân độ trễ.")]
    for index,(name,caption) in enumerate(figures): blocks.append({"kind":"landscape_figure","path":"figures/"+name+".png","caption":caption,"appendixStart":index==0})
    return blocks,references


def font(run,size=None,bold=None):
    run.font.name="Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"),"Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman")
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)


def build_docx(blocks,path,authors):
    doc=Document()
    section=doc.sections[0]
    def setup(s,landscape=False):
        s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        s.page_width=Cm(29.7 if landscape else 21);s.page_height=Cm(21 if landscape else 29.7)
        s.top_margin=Cm(1.7 if landscape else 2.0);s.bottom_margin=Cm(1.6 if landscape else 2.0)
        s.left_margin=Cm(2.2);s.right_margin=Cm(2.2);s.footer_distance=Cm(.8)
    setup(section)
    for name in ("Normal","Title","Subtitle","Heading 1","Heading 2","Caption"):
        style=doc.styles[name]
        style.font.name="Times New Roman";style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt({"Title":25,"Subtitle":14,"Heading 1":15,"Heading 2":12.5,"Caption":10.5}.get(name,12))
        pf=style.paragraph_format
        pf.space_after=Pt(6);pf.line_spacing=1.13
        if name.startswith("Heading"):pf.space_before=Pt(9);pf.keep_with_next=True
        rfonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"):
            rfonts.attrib.pop(qn("w:"+key),None)
        rfonts.set(qn("w:ascii"),"Times New Roman");rfonts.set(qn("w:hAnsi"),"Times New Roman")
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.styles["Title"].paragraph_format.space_after=Pt(6)
    doc.styles["Title"].paragraph_format.keep_with_next=True
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Trang "),9)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
    field_run=OxmlElement("w:r");rpr=OxmlElement("w:rPr");sz=OxmlElement("w:sz");sz.set(qn("w:val"),"18");rpr.append(sz);field_run.append(rpr)
    text=OxmlElement("w:t");text.text="1";field_run.append(text);field.append(field_run);footer._p.append(field)
    for block in blocks:
        kind=block["kind"]
        if kind=="page":doc.add_page_break()
        elif kind=="heading":doc.add_heading(block["text"],level=block["level"])
        elif kind in ("title","subtitle"):
            p=doc.add_paragraph(style=kind.title());font(p.add_run(block["text"]),25 if kind=="title" else 14,kind=="title")
        elif kind=="paragraph":
            p=doc.add_paragraph();p.paragraph_format.widow_control=True
            font(p.add_run(block["text"]))
        elif kind=="code":
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.0
            r=p.add_run(block["text"]);r.font.name="Consolas";r.font.size=Pt(9.5)
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"F3F5F7");p._p.get_or_add_pPr().append(shade)
        elif kind=="table":
            p=doc.add_paragraph(style="Caption");p.paragraph_format.keep_with_next=True;font(p.add_run(block["caption"]),10.5,True)
            t=doc.add_table(rows=1,cols=len(block["headers"]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            widths=block.get("widths") or [16.6/len(block["headers"])]*len(block["headers"])
            total=sum(widths);widths=[x*16.6/total for x in widths]
            for cell,width in zip(t.rows[0].cells,widths):cell.width=Cm(width)
            for cell,text in zip(t.rows[0].cells,block["headers"]):cell.text=str(text)
            header=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(header)
            for row in block["rows"]:
                cells=t.add_row().cells
                for cell,value,width in zip(cells,row,widths):cell.text=str(value);cell.width=Cm(width)
            for i,row in enumerate(t.rows):
                trpr=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");trpr.append(cant)
                for j,cell in enumerate(row.cells):
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    shd=OxmlElement("w:shd");shd.set(qn("w:fill"),"DCE6EF" if i==0 else ("F5F7F9" if i%2==0 else "FFFFFF"));tcpr.append(shd)
                    margins=OxmlElement("w:tcMar")
                    for edge in ("top","bottom","left","right"):
                        e=OxmlElement("w:"+edge);e.set(qn("w:w"),"75");e.set(qn("w:type"),"dxa");margins.append(e)
                    tcpr.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(block["headers"])>=6 or j==0 else WD_ALIGN_PARAGRAPH.LEFT
                        for r in p.runs:font(r,10 if len(block["headers"])>=6 else 10.5,i==0)
            borders=OxmlElement("w:tblBorders")
            for edge in ("top","left","bottom","right","insideH","insideV"):
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
            t._tbl.tblPr.append(borders)
            gap=doc.add_paragraph();gap.paragraph_format.space_after=Pt(1);gap.paragraph_format.space_before=Pt(0);font(gap.add_run(""),2)
        elif kind=="figure":
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(REPORT/block["path"]),width=Cm(16.3))
            p.runs[-1]._r.xpath('.//wp:docPr')[0].set('descr',block['caption'])
            p=doc.add_paragraph(style="Caption");font(p.add_run(block["caption"]),10.5)
        elif kind=="landscape_figure":
            s=doc.add_section(WD_SECTION_START.NEW_PAGE);setup(s,True)
            if block.get('appendixStart'):doc.add_heading('Phụ lục B Hình thực nghiệm',level=1)
            p=doc.add_paragraph(style="Caption");p.paragraph_format.keep_with_next=True;font(p.add_run(block["caption"]),11)
            with Image.open(REPORT/block["path"]) as im:w,h=im.size
            width=min(25.3,14.8*w/h)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(REPORT/block["path"]),width=Cm(width))
            p.runs[-1]._r.xpath('.//wp:docPr')[0].set('descr',block['caption'])
    doc.core_properties.title="Huyền Cơ ứng dụng tra cứu tri thức qua mạng"
    doc.core_properties.subject="Báo cáo bài tập lớn Lập trình mạng và thực nghiệm X01"
    doc.core_properties.author="; ".join(m['name'] for m in authors.get('members',[])) or "Nhóm phát triển Huyền Cơ"
    doc.core_properties.keywords="WebSocket, STOMP, bounded queue, idempotency, provenance"
    doc.core_properties.comments=""
    settings=doc.settings.element
    update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)
    for border in doc.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    staging=path.with_name(path.stem+'.building.docx')
    doc.save(staging)
    os.replace(staging,path)


def markdown(blocks):
    parts=[]
    for b in blocks:
        if b["kind"] in ("heading","title"):
            parts.append("#"*(b.get("level",0)+1)+" "+b["text"])
        elif b["kind"] in ("paragraph","subtitle"):parts.append(b["text"])
        elif b["kind"]=="code":parts.append("```json\n"+b["text"]+"\n```")
        elif b["kind"]=="table":
            parts.append(b["caption"])
            parts.append("| "+" | ".join(b["headers"])+" |\n| "+" | ".join(["---"]*len(b["headers"]))+" |\n"+"\n".join("| "+" | ".join(str(x) for x in r)+" |" for r in b["rows"]))
        elif b["kind"] in ("figure","landscape_figure"):parts.append("!["+b["caption"]+"]("+b["path"]+")")
    return "\n\n".join(parts)+"\n"


def main():
    checkpoint,analysis,publication=load_evidence()
    REPORT.mkdir(exist_ok=True);(REPORT/"figures").mkdir(exist_ok=True)
    metadata=REPORT/"authors.json"
    if not metadata.exists():metadata.write_text(json.dumps({"groupNumber":None,"members":[]},ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    authors=read(metadata)
    group=authors.get("groupNumber")
    if group is not None and (not str(group).isdigit() or not 1<=int(group)<=99):raise ValueError("Group number must be 1–99")
    name=(f"Nhom{int(group):02d}_" if group is not None else "")+"HuyenCo_Report.docx"
    architecture(REPORT/"figures/architecture.png")
    # Preserve axes, samples and full distributions. Only duplicate title/footer space is cropped.
    for figure_name in analysis["figures"]:
        source=ROOT/"docs/assets/x01"/(figure_name+".png")
        with Image.open(source) as im:
            im.crop((0,155,im.width,im.height-(150 if figure_name=="04-resources" else 125))).save(REPORT/"figures"/(figure_name+".png"))
    blocks,references=content(analysis,authors)
    build_docx(blocks,REPORT/name,authors)
    (REPORT/"HuyenCo_Report.md").write_text(markdown(blocks),encoding="utf-8")
    evidence={"task":"X01.7","status":"CONTENT_READY_PENDING_RENDER_REVIEW","createdAt":datetime.now(VN).isoformat(),
        "report":str((REPORT/name).relative_to(ROOT)).replace("\\","/"),"analysisCheckpoint":checkpoint,
        "analysisPublicationSha256":digest(ROOT/"docs/assets/x01/publication-manifest.json"),
        "analysisReportSha256":digest(ROOT/"docs/53-X01-BENCHMARK-ANALYSIS.md"),
        "matrixAuditStatus":analysis["matrixAuditStatus"],"strictSendWindowSatisfied":False,
        "groupMetadataComplete":bool(group and authors.get("members")),"compilatioRun":False,"submitted":False,
        "figures":{p.name:digest(p) for p in sorted((REPORT/"figures").glob("*.png"))},
        "references":[{"number":n,"text":text} for n,text in references],
        "sourceFiles":{str(p.relative_to(ROOT)).replace("\\","/"):digest(p) for p in [ROOT/"docs/assets/x01/per-run.csv",ROOT/"docs/assets/x01/aggregate.csv",ROOT/"docs/assets/x01/boundary-exclusions.csv",ROOT/"scripts/build_x01_report.py"]}}
    (REPORT/"evidence-index.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"report":str(REPORT/name),"blocks":len(blocks),"groupMetadataComplete":evidence["groupMetadataComplete"]},ensure_ascii=False))


if __name__=="__main__":main()
